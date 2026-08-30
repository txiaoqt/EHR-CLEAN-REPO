# src/scripts/test_case_generator/generator_core.py
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

class TestCase:
    def __init__(self, id, description, preconditions, steps, test_data, expected_result, pass_criteria, fail_criteria):
        self.id = id
        self.description = description
        self.preconditions = preconditions
        self.steps = steps if isinstance(steps, list) else [steps]
        self.test_data = test_data
        self.expected_result = expected_result
        self.actual_result = ""  # MUST BE BLANK for manual tester entry
        self.pass_criteria = pass_criteria
        self.fail_criteria = fail_criteria

class TestSection:
    def __init__(self, number, title, description, test_cases=None):
        self.number = number
        self.title = title
        self.description = description
        self.test_cases = test_cases or []

def set_cell_background(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_styled_document():
    doc = docx.Document()
    
    # Page setup - Normal 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    # Configure default style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Segoe UI'
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(0x1F, 0x29, 0x37) # Slate 800
    
    return doc

def add_header_banner(doc, total_cases, module_count):
    # Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("TUP CLINIC EHR SYSTEM")
    title_run.font.name = 'Segoe UI'
    title_run.font.size = Pt(24)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0x99, 0x1B, 0x1B) # TUP Crimson Maroon
    
    subtitle_p = doc.add_paragraph()
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = subtitle_p.add_run("COMPREHENSIVE BLACK-BOX TEST SUITE & QUALITY ASSURANCE SPECIFICATION")
    sub_run.font.name = 'Segoe UI'
    sub_run.font.size = Pt(13)
    sub_run.font.bold = True
    sub_run.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
    
    doc.add_paragraph()
    
    # Metadata Box Table
    meta_table = doc.add_table(rows=6, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    
    meta_data = [
        ("System / Application:", "TUP Clinic Electronic Health Records (EHR) System"),
        ("Testing Methodology:", "Strict Black-Box Functional & End-to-End QA Testing"),
        ("Deployment Surfaces:", "Student/Patient Portal & Clinical Administrative Staff Portal"),
        ("User Roles Covered:", "Patient (Student), Nurse, Physician, Administrator"),
        ("Total Functional Modules:", f"{module_count} Modules & Test Sections"),
        ("Total Black-Box Test Cases:", f"{total_cases} Globally Sequential Test Cases (TC001 – TC{total_cases:03d})"),
    ]
    
    for i, (label, val) in enumerate(meta_data):
        row = meta_table.rows[i]
        c0 = row.cells[0]
        c1 = row.cells[1]
        c0.width = Inches(2.3)
        c1.width = Inches(4.2)
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(label)
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(val)
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
        
        set_cell_background(c0, "F3F4F6")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, top=60, bottom=60, left=100, right=100)
        set_cell_margins(c1, top=60, bottom=60, left=100, right=100)
        
    doc.add_paragraph()
    doc.add_page_break()

def add_table_of_contents(doc, sections):
    toc_heading = doc.add_paragraph()
    r = toc_heading.add_run("Table of Contents")
    r.font.name = 'Segoe UI'
    r.font.size = Pt(18)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x99, 0x1B, 0x1B)
    
    toc_desc = doc.add_paragraph()
    r_desc = toc_desc.add_run("Directory of testing sections, modules, and test case ID ranges:")
    r_desc.font.italic = True
    r_desc.font.size = Pt(10)
    
    toc_table = doc.add_table(rows=len(sections) + 1, cols=4)
    toc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    toc_table.autofit = False
    
    headers = ["Sec #", "Module / Section Title", "Test ID Range", "Case Count"]
    widths = [Inches(0.6), Inches(3.6), Inches(1.3), Inches(1.0)]
    
    hdr_row = toc_table.rows[0]
    for j, h in enumerate(headers):
        cell = hdr_row.cells[j]
        cell.width = widths[j]
        p = cell.paragraphs[0]
        run = p.add_run(h)
        run.font.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_background(cell, "991B1B")
        set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
        
    for i, sec in enumerate(sections):
        row = toc_table.rows[i + 1]
        start_id = sec.test_cases[0].id if sec.test_cases else "N/A"
        end_id = sec.test_cases[-1].id if sec.test_cases else "N/A"
        id_range = f"{start_id} – {end_id}" if start_id != end_id else start_id
        
        vals = [f"{sec.number}", sec.title, id_range, str(len(sec.test_cases))]
        bg_color = "F9FAFB" if i % 2 == 0 else "FFFFFF"
        
        for j, val in enumerate(vals):
            cell = row.cells[j]
            cell.width = widths[j]
            p = cell.paragraphs[0]
            if j in [0, 2, 3]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(val)
            run.font.size = Pt(9)
            if j == 1:
                run.font.bold = True
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            
    doc.add_page_break()

def render_test_case_table(doc, tc):
    tbl = doc.add_table(rows=9, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    w_label = Inches(1.6)
    w_val = Inches(4.9)
    
    # 0: Header Banner
    header_row = tbl.rows[0]
    header_cell = header_row.cells[0]
    header_cell.merge(header_row.cells[1])
    p_hdr = header_cell.paragraphs[0]
    r_hdr = p_hdr.add_run(f"[{tc.id}] {tc.description}")
    r_hdr.font.bold = True
    r_hdr.font.size = Pt(10.5)
    r_hdr.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    set_cell_background(header_cell, "1E3A8A") # Deep Navy Blue
    set_cell_margins(header_cell, top=70, bottom=70, left=100, right=100)
    
    fields = [
        ("Test Case ID", tc.id),
        ("Description", tc.description),
        ("Preconditions", tc.preconditions),
        ("Test Steps", "\n".join([f"{idx+1}. {s}" for idx, s in enumerate(tc.steps)])),
        ("Test Data", tc.test_data),
        ("Expected Result", tc.expected_result),
        ("Actual Result", ""), # Left completely blank for manual tester entry
    ]
    
    for row_idx, (label, val) in enumerate(fields, start=1):
        row = tbl.rows[row_idx]
        c_label = row.cells[0]
        c_val = row.cells[1]
        c_label.width = w_label
        c_val.width = w_val
        
        p_lbl = c_label.paragraphs[0]
        r_lbl = p_lbl.add_run(label)
        r_lbl.font.bold = True
        r_lbl.font.size = Pt(9)
        r_lbl.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
        
        p_val = c_val.paragraphs[0]
        r_val = p_val.add_run(val)
        r_val.font.size = Pt(9)
        r_val.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
        
        set_cell_background(c_label, "F3F4F6")
        set_cell_background(c_val, "FFFFFF")
        
        # Give Actual Result dedicated blank height for tester entry
        if label == "Actual Result":
            set_cell_margins(c_label, top=100, bottom=100, left=80, right=80)
            set_cell_margins(c_val, top=140, bottom=140, left=80, right=80)
        else:
            set_cell_margins(c_label, top=50, bottom=50, left=80, right=80)
            set_cell_margins(c_val, top=50, bottom=50, left=80, right=80)
        
    # Pass/Fail Criteria in final row (row 8)
    crit_row = tbl.rows[8]
    c_crit_lbl = crit_row.cells[0]
    c_crit_val = crit_row.cells[1]
    c_crit_lbl.width = w_label
    c_crit_val.width = w_val
    
    p_c_lbl = c_crit_lbl.paragraphs[0]
    r_c_lbl = p_c_lbl.add_run("Pass/Fail Criteria")
    r_c_lbl.font.bold = True
    r_c_lbl.font.size = Pt(9)
    r_c_lbl.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
    
    p_c_val = c_crit_val.paragraphs[0]
    r_pass = p_c_val.add_run(f"Pass:\n{tc.pass_criteria}\n\n")
    r_pass.font.size = Pt(9)
    r_pass.font.color.rgb = RGBColor(0x06, 0x5F, 0x46) # Emerald text
    
    r_fail = p_c_val.add_run(f"Fail:\n{tc.fail_criteria}")
    r_fail.font.size = Pt(9)
    r_fail.font.color.rgb = RGBColor(0x99, 0x1B, 0x1B) # Crimson text
    
    set_cell_background(c_crit_lbl, "F3F4F6")
    set_cell_background(c_crit_val, "FAFAFA")
    set_cell_margins(c_crit_lbl, top=50, bottom=50, left=80, right=80)
    set_cell_margins(c_crit_val, top=50, bottom=50, left=80, right=80)
    
    doc.add_paragraph() # Spacing between test cases
